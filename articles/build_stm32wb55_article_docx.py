# -*- coding: utf-8 -*-
"""Двухстраничная статья: защита прошивки контроллера протеза на STM32WB55."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

ROOT = Path(__file__).resolve().parent
DOCX = ROOT / "stm32wb55_prosthesis_protection.docx"
MD = ROOT / "stm32wb55_prosthesis_protection.md"
FONT = "Times New Roman"

TITLE = (
    "Защита прошивки беспроводного контроллера протеза "
    "на микроконтроллере STM32WB55 от аппаратного реверс-инжиниринга"
)

AUTHOR = "Ф. И. О. автора, НИЯУ МИФИ, г. Москва"

ABSTRACT = (
    "Рассмотрена защита встроенного программного обеспечения контроллера протеза "
    "на STM32WB55 при физическом доступе к плате. Показано, что основной путь "
    "реверс-инжиниринга — открытый отладчик и незащищённое обновление по Bluetooth "
    "Low Energy. Температурное воздействие и анализ профиля энергопотребления "
    "становятся значимыми после закрытия чтения Flash. Для серии WB зафиксированы "
    "доступные средства: уровни запрета чтения RDP, область PCROP, хранилище ключей "
    "FUS, датчик температуры АЦП, детектор напряжения и внешние входы tamper. "
    "Предложен порядок мер. Необратимое отключение отладчика допустимо только при "
    "уже действующем подписанном канале обновления."
)

KEYWORDS = (
    "STM32WB55, протез, реверс-инжиниринг, RDP, PCROP, анализ энергопотребления, "
    "температурный сбой, Bluetooth Low Energy"
)

SECTIONS = [
    (
        "Введение",
        [
            "Контроллер протеза хранит закон управления хватом или движением, "
            "обработку сигналов мышц и ключи радиоканала, по которому приходит "
            "обновление прошивки. Плата доступна при сервисе, утере или разборке "
            "изделия, поэтому угроза не сводится к удалённой атаке на протокол. "
            "Для микроконтроллеров STM32 чтение встроенной Flash через отладочный "
            "интерфейс остаётся первым и самым дешёвым способом получить прошивку [1]. "
            "Нагрев кристалла и съём тока питания относятся к следующему слою: они "
            "нужны, когда прямое чтение памяти уже закрыто, и направлены на ключ "
            "шифрования либо на пропуск проверки подписи [2–4].",
            "Цель работы — связать модель физического доступа к контроллеру протеза "
            "на STM32WB55 с механизмами защиты этой серии и задать порядок их включения. "
            "Работа не переносит на WB55 приёмы, показанные для других кристаллов, "
            "и не описывает процедуру вскрытия.",
        ],
    ),
    (
        "Модель угрозы",
        [
            "Реалистичная последовательность состоит из четырёх шагов. Сначала читают "
            "Flash через SWD, если вывод отладчика оставлен на разъёме, а защита чтения "
            "выключена. Затем забирают образ из эфира, если обновление по BLE передаётся "
            "без подписи. Если оба пути закрыты, по падению напряжения на шунте питания "
            "восстанавливают ключ программного AES: для Cortex-M показано, что аппаратный "
            "блок AES без дополнительных мер отдаёт ключ по профилю потребления, а "
            "случайная задержка перед операцией отфильтровывается обработкой трасс [4]. "
            "Наконец, при повышенной температуре расширяется интервал, в котором короткий "
            "сбой такта пропускает инструкцию [2, 3]. Для загрузчика это пропуск проверки "
            "подписи.",
            "Серия WB55 не повторяет STM32F0, на которой ранее показана недостаточность "
            "одного только уровня запрета чтения [1]. Аппаратного внутреннего "
            "температурного tamper, как у старших линеек, у WB55 нет. Есть датчик "
            "температуры на АЦП, детектор напряжения и внешние входы tamper часов "
            "реального времени [5].",
        ],
    ),
    (
        "Меры на STM32WB55",
        [
            "Уровень RDP 1 запрещает доступ к Flash при подключённом отладчике; снятие "
            "уровня стирает память. Уровень 2 отключает отладчик необратимо [5, 6]. "
            "На опытных протезах достаточен уровень 1: ошибку закона управления ещё "
            "можно исправить. Уровень 2 включают на экземпляре, у которого уже работает "
            "подписанное обновление, иначе изделие нельзя обслуживать.",
            "Закон управления и обработку миосигнала помещают в область PCROP: код из "
            "неё исполняется, но не читается ни отладчиком, ни остальной программой [6]. "
            "Ключ подписи образа не хранят массивом во Flash ядра Cortex-M4. Его помещают "
            "в хранилище FUS: загрузка в AES выполняется со стороны Cortex-M0+, и "
            "прикладное ядро регистр ключа не видит [5]. Обновление по BLE принимается "
            "только при верной подписи. Проверку выполняют дважды, разными участками "
            "кода. Простое дублирование одной инструкции не держит сбой, если пропускается "
            "больше одной команды, и усиливает утечку по питанию [7].",
            "Перед приёмом прошивки и перед разрешением движения читают датчик "
            "температуры и порог напряжения. Выход за рабочий диапазон блокирует "
            "обновление и привод. Случайная задержка перед проверкой подписи затрудняет "
            "совмещение трасс, но сама по себе прошивку не скрывает [4, 8]. Коэффициенты "
            "регулятора не публикуют открытыми характеристиками BLE.",
        ],
    ),
]

CONCLUSION = (
    "Для контроллера протеза на STM32WB55 реверс-инжиниринг закрывается в первую "
    "очередь запретом чтения памяти, подписью обновления и выносом ключа и закона "
    "управления из доступной области Flash. Анализ питания и температурный сбой "
    "задают требования к загрузчику: не считать AES на прикладном ядре с долгоживущим "
    "ключом и не продолжать обновление или движение вне допустимой температуры и "
    "напряжения. Необратимая блокировка отладчика уместна только после появления "
    "рабочего канала обслуживания."
)

REFS = [
    "Obermaier J., Tatschner S. Shedding too much Light on a Microcontroller's Firmware Protection // 11th USENIX Workshop on Offensive Technologies. 2017. URL: https://www.usenix.org/system/files/conference/woot17/woot17-paper-obermaier.pdf (дата обращения: 02.10.2026).",
    "Hutter M., Schmidt J.-M. The Temperature Side Channel and Heating Fault Attacks // CARDIS 2013. URL: https://eprint.iacr.org/2014/190 (дата обращения: 02.10.2026).",
    "Korak T., Hutter M., Ege B., Batina L. Clock Glitch Attacks in the Presence of Heating // FDTC 2014. DOI: 10.1109/FDTC.2014.19.",
    "Unterstein F., Schink M., Schamberger T., Tebelmann L., Ilg M., Heyszl J. Retrofitting Leakage Resilient Authenticated Encryption to Microcontrollers // IACR Transactions on Cryptographic Hardware and Embedded Systems. 2020. URL: https://eprint.iacr.org/2020/960 (дата обращения: 02.10.2026).",
    "STMicroelectronics. Introduction to security for STM32 MCUs : AN5156, rev. 13. 2026. URL: https://www.st.com/resource/en/application_note/an5156-introduction-to-security-for-stm32-mcus-stmicroelectronics.pdf (дата обращения: 02.10.2026).",
    "STMicroelectronics. Introduction to proprietary code read-out protection for STM32L4, STM32L4+, STM32G4 and STM32WB MCUs : AN4758, rev. 7. 2026.",
    "Cojocar L., Papagiannopoulos K., Timmers N. Instruction Duplication: Leaky and Not Too Fault-Tolerant! // CARDIS 2017. URL: https://eprint.iacr.org/2017/1082 (дата обращения: 02.10.2026).",
    "Leplus G., Savry O., Bossuet L. Insertion of random delay with context-aware dummy instructions generator in a RISC-V processor // IEEE HOST 2022. URL: https://hal.science/hal-04004056 (дата обращения: 02.10.2026).",
]


def _font(run, size=14, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), FONT)


def _fmt(p, *, first=True, align="justify", before=0, after=0, spacing=1.15, size_indent=True):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.first_line_indent = Cm(1.0) if first else Cm(0)
    pf.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]


def add_p(doc, text, *, first=True, align="justify", size=14, bold=False, italic=False,
          before=0, after=0, spacing=1.15, keep_with_next=False):
    p = doc.add_paragraph()
    _fmt(p, first=first, align=align, before=before, after=after, spacing=spacing)
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    _font(run, size=size, bold=bold, italic=italic)
    return p


def write_md():
    lines = [
        f"# {TITLE}",
        "",
        AUTHOR,
        "",
        f"**Аннотация.** {ABSTRACT}",
        "",
        f"**Ключевые слова:** {KEYWORDS}",
        "",
    ]
    for title, paragraphs in SECTIONS:
        lines.append(f"## {title}")
        lines.append("")
        for paragraph in paragraphs:
            lines.append(paragraph)
            lines.append("")
    lines.append("## Заключение")
    lines.append("")
    lines.append(CONCLUSION)
    lines.append("")
    lines.append("## Список литературы")
    lines.append("")
    for i, ref in enumerate(REFS, start=1):
        lines.append(f"{i}. {ref}")
        lines.append("")
    MD.write_text("\n".join(lines), encoding="utf-8")


def write_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(1.5)

    add_p(doc, TITLE, first=False, align="center", bold=True, before=0, after=6, spacing=1.0)
    add_p(doc, AUTHOR, first=False, align="center", italic=True, size=12, after=6, spacing=1.0)
    add_p(doc, "Аннотация. " + ABSTRACT, first=True, size=12, spacing=1.0, after=2)
    add_p(doc, "Ключевые слова: " + KEYWORDS, first=True, size=12, italic=True, spacing=1.0, after=6)

    for title, paragraphs in SECTIONS:
        add_p(doc, title, first=False, align="center", bold=True, before=6, after=2, spacing=1.0,
              keep_with_next=True)
        for paragraph in paragraphs:
            add_p(doc, paragraph, spacing=1.0, after=2)

    add_p(doc, "Заключение", first=False, align="center", bold=True, before=6, after=2, spacing=1.0,
          keep_with_next=True)
    add_p(doc, CONCLUSION, spacing=1.0, after=4)
    add_p(doc, "Список литературы", first=False, align="center", bold=True, before=4, after=2, spacing=1.0,
          keep_with_next=True)
    for i, ref in enumerate(REFS, start=1):
        add_p(doc, f"{i}. {ref}", first=False, size=11, spacing=1.0, after=1)

    doc.save(DOCX)


def main():
    write_md()
    write_docx()
    text = MD.read_text(encoding="utf-8")
    print(f"chars={len(text)} md={MD.name} docx={DOCX.name}")


if __name__ == "__main__":
    main()
